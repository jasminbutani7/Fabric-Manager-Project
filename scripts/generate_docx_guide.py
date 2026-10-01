import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = Document()

# Set clean margins
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def add_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(30, 41, 59) # Slate 800
    p.paragraph_format.space_after = Pt(2)

def add_subtitle(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(12)
    run.font.italic = True
    run.font.color.rgb = RGBColor(79, 70, 229) # Indigo 600
    p.paragraph_format.space_after = Pt(14)

def add_h1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42)

def add_h2(text, color=(67, 56, 202)):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(*color)

def add_p(text, bold=False, italic=False, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(10)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_bullet(text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(51, 65, 85)

def add_callout(title, text, bg="EEF2FF", border="4338CA"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, bg)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    
    r_title = p.add_run(title + " ")
    r_title.font.name = 'Arial'
    r_title.font.bold = True
    r_title.font.size = Pt(10)
    r_title.font.color.rgb = RGBColor(67, 56, 202)
    
    r_text = p.add_run(text)
    r_text.font.name = 'Arial'
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

# DOCUMENT HEADER
add_title("Fabric Manager - Step-by-Step Beginner Guide")
add_subtitle("Simple, Practical Guide to Build the Entire Web App from Scratch")

add_callout("TEAM ASSIGNMENT:", "Lead: Sharad | Frontend: Nandni & Tisha | Backend: Sharad & Jasmin")

add_h1("💡 What is this Project About? (In Simple Words)")
add_p("Imagine a garment factory in real life:")
add_bullet("1. The factory buys big rolls of cloth (called Fabric Lots).")
add_bullet("2. A Cutting Master cuts the cloth into shirt or pant pieces (Front, Back, Sleeves).")
add_bullet("3. Tailors sit at sewing machines and stitch the pieces together into finished clothes.")
add_bullet("4. Helpers trim the extra threads and pack the dresses.")
add_bullet("5. The factory delivers the finished clothes to the buyer and makes a Bill (Invoice).")
add_bullet("6. The factory owner pays worker wages and checks daily profit.")
add_p("Our app, Fabric Manager, does all of this on a mobile phone and computer!")

# PART 1: FRONTEND
add_h1("🎨 PART 1: FRONTEND STEP-BY-STEP GUIDE (Nandni & Tisha)")
add_p("Build the application step-by-step in this exact sequence:")

add_h2("Step 1: Admin Signup & Login Screens (SignUp.jsx & SignIn.jsx)")
add_p("The app must first allow factory owners to create an account and sign in.")
add_bullet("Signup Page (/signup): Create a clean form with Factory Name, Admin Name, Mobile Number, Email, and Password. Add a green/indigo 'Create Account' button.")
add_bullet("Login Page (/signin): Create form with Email/Phone and Password. When logged in, save the session in sessionStorage so the user stays logged in.")
add_bullet("Route Protection (App.jsx): If the user is NOT logged in, redirect them to /signin. If logged in, open the Home dashboard.")

add_h2("Step 2: Main App Shell & Home Page (MobileLayout.jsx & Home.jsx)")
add_p("Create the look and feel of the app with mobile-friendly navigation.")
add_bullet("Top Header Bar: Display factory name (e.g. 'Krishna Textiles') with a small green pulse light (Live indicator) and a menu button.")
add_bullet("Bottom Navigation Bar: 5 primary tabs with icons: 🏠 Home, 🧵 Fabrics, ✂️ Production, 💳 Transaction, ⚙️ Setup.")
add_bullet("Home Screen Content (Home.jsx): Show active machines count, tailors at work, and running cloth lots with quick action buttons (+ Add Cloth, + Cut Cloth).")

add_h2("Step 3: Setup Page - Workers & Machines (Setup.jsx)")
add_p("Before cutting cloth, we need machines and workers registered in the app.")
add_bullet("Tab Bar at Top: Machines | Tailors | Cutting Masters | Helpers | Suppliers.")
add_bullet("Add Machine Modal: Enter machine code (e.g. M-01), type (Single needle, Overlock), and brand (Juki/Jack).")
add_bullet("Add Tailor Modal: Enter tailor name, mobile, and assign them to a machine.")
add_bullet("Add Master Modal: Enter Cutting Master name and years of experience.")
add_bullet("Add Helper Modal: Enter Helper name and default rate per piece (e.g. Rs. 2/pc).")
add_bullet("Add Supplier Modal: Enter fabric mill name, contact phone, and GST number.")

add_h2("Step 4: Fabric Inventory Page (Inventory.jsx)")
add_p("Record cloth arriving at the factory and configure dress structures.")
add_bullet("Add Fabric Lot Modal: Enter Lot Number (e.g. Lot #101), select Supplier, enter cloth plain meters (e.g. 500m), color, and price per meter.")
add_bullet("Product Structure Modal: Build dress structure: Top Kurti (Rate Rs. 25), Pant (Rate Rs. 15), Dupatta (Rate Rs. 5). Under each, list components (Front, Back, Sleeves).")
add_bullet("Lot Cards: Show lot card with 5-step progress line (1. Inward -> 2. Cutting -> 3. Stitching -> 4. Handover -> 5. Delivery).")

add_h2("Step 5: Production & Cutting Execution (Orders.jsx)")
add_p("Cutting master cuts cloth, and pieces are given to tailors.")
add_bullet("Record Cutting Modal: Select Cutting Master, enter meters of cloth cut (e.g. 120m), enter pieces cut (50 Front, 50 Back, 100 Sleeves). App automatically calculates ready complete sets and remaining cloth meters!")
add_bullet("Assign Tailors Modal: Pick cut pieces, select a tailor, enter quantity to stitch (e.g. 50 Tops). App automatically shows total rupees the tailor will earn (50 x Rs. 25 = Rs. 1,250).")

add_h2("Step 6: Dedicated Worker Mobile Portals (Self-Service)")
add_p("Workers open their own portal link on their mobile without seeing admin data.")
add_bullet("Tailor Portal (/portal/tailor/:id): Clean mobile screen for tailors. Card says 'New Job Assigned: 50 Tops at Rs. 25/pc'. Green 'Accept Job' button. 'Log Today's Work' popup to type pieces stitched today. Shows total earned money in Rupees.")
add_bullet("Cutting Master Portal (/portal/master/:id): Screen for Cutting Master to see assigned rolls. 1-tap 'Approve' button to check tailor submissions. 'Handover to Admin' popup to send finished goods back to manager.")

add_h2("Step 7: Deliveries, Billing & Expenses (Transactions.jsx)")
add_p("Deliver finished clothes to buyer, create invoices, and log factory expenses.")
add_bullet("Thread Trimming Popup: Ask admin: 'Does this lot need thread cutting?' If Yes, assign helper at Rs. 2/pc.")
add_bullet("Deliver to Buyer Modal: Enter vehicle number, delivery date, and quantity of dresses dispatched.")
add_bullet("Create Invoice Modal: Select buyer, enter rate per dress, add GST Tax % and discount. Calculates grand total.")
add_bullet("Record Payment Modal: If client gives partial payment (e.g. Rs. 10,000 out of Rs. 30,000), record payment. Updates status to 'Partial' and shows remaining due balance.")
add_bullet("Factory Expenses (Expenses.jsx): Record bills for Rent, Electricity, Machine Needles, Thread, and Tea/Snacks.")

# PART 2: BACKEND
add_h1("⚙️ PART 2: BACKEND STEP-BY-STEP GUIDE (Jasmin & Sharad)")
add_p("Build the database and data connections step-by-step in this exact sequence:")

add_h2("Step 1: Create Supabase Project & Connect App")
add_p("Connect the React frontend to the PostgreSQL cloud database.")
add_bullet("Create free project on Supabase (supabase.com).")
add_bullet("Get Project URL and Anon Public Key, and put them in .env file.")
add_bullet("Create src/lib/supabaseClient.js to initialize Supabase.")

add_h2("Step 2: Create Users Table (For Admin Logins)")
add_p("Table users: id (text), name, email, phone, company_name, password, created_at.")

add_h2("Step 3: Create Master Data Tables")
add_bullet("Table masters: id, user_id, name, phone, experience, status.")
add_bullet("Table machines: id, user_id, code, name, type, brand, assigned_tailor_id, status.")
add_bullet("Table tailors: id, user_id, name, phone, assigned_machine_id, specialty, status.")
add_bullet("Table helpers: id, user_id, name, phone, default_rate_per_piece, status.")
add_bullet("Table suppliers: id, user_id, name, phone, gst_number, payment_terms, status.")

add_h2("Step 4: Create Fabric & Cutting Tables")
add_bullet("Table fabrics: id, user_id, lot_no, supplier_id, plain_meters, target_sets, price_per_set, product_name, product_structure (JSONB), status.")
add_bullet("Table cuttings: id, user_id, fabric_id, master_id, cut_pieces (JSONB), total_pieces, sets_count, meters_used, status.")

add_h2("Step 5: Create Stitching & Handover Tables")
add_bullet("Table tailor_assignments: id, user_id, cutting_id, tailor_id, assigned_cut_pieces (JSONB), rate_per_piece, completed_quantity, total_earned, daily_submissions (JSONB), tailor_accepted (boolean), status.")
add_bullet("Table master_handovers: id, user_id, fabric_id, cutting_id, master_id, pieces_count, status ('Pending', 'Verified').")

add_h2("Step 6: Create Billing & Expense Tables")
add_bullet("Table delivery_logs: id, user_id, fabric_id, supplier_id, total_pieces, vehicle_no, delivery_date.")
add_bullet("Table invoices: id, user_id, invoice_number, client_name, date, due_date, subtotal, tax, discount, total_amount, paid_amount, due_amount, payments (JSONB), status ('Pending', 'Partial', 'Paid').")
add_bullet("Table expenses: id, user_id, category, title, amount, date, payment_method, paid_to.")

add_h2("Step 7: Turn On Realtime Live Sync (Important!)")
add_p("Enable Realtime in Supabase so that when a tailor logs work on phone, the admin dashboard updates instantly without refreshing! Enable for: fabrics, cuttings, tailor_assignments, master_handovers, invoices.")

# TIMELINE
add_h1("🚀 4-Week Simple Timeline for Team")
time_table = doc.add_table(rows=1, cols=4)
time_table.alignment = WD_TABLE_ALIGNMENT.CENTER
thdr = time_table.rows[0].cells
thdr[0].text = "Week"
thdr[1].text = "Nandni & Tisha (Frontend)"
thdr[2].text = "Jasmin & Sharad (Backend)"
thdr[3].text = "Weekly Result"
for c in thdr:
    set_cell_background(c, "1E293B")
    for p in c.paragraphs:
        for r in p.runs:
            r.font.name = 'Arial'
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

trows = [
    ("Week 1", "Build Login, Signup, Layout & Setup Page (Workers + Machines)", "Create Supabase project, Setup users, machines, tailors tables", "Login works, and you can register tailors & machines!"),
    ("Week 2", "Build Fabric Inventory screen & Cutting Batch popup", "Create fabrics and cuttings tables with remaining meters calculation", "You can add cloth lots and record cuts!"),
    ("Week 3", "Build Tailor Portal on mobile & Daily Work logging popup", "Create tailor_assignments table & enable Realtime live sync", "Tailor can accept jobs and log work from phone!"),
    ("Week 4", "Build Deliveries, Invoices (Billing) & Expense screen", "Create invoices, delivery_logs, and expenses tables", "Complete app ready with live billing & profit tracking!")
]

for row_data in trows:
    row = time_table.add_row()
    for i, text in enumerate(row_data):
        cell = row.cells[i]
        cell.text = text
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.name = 'Arial'
                run.font.size = Pt(9)

output_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), "Fabric_Manager_Team_Project_Guide.docx")
doc.save(output_path)
print(f"User-friendly Docx file generated successfully at: {output_path}")
